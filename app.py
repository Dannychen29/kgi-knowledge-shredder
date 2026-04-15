from flask import Flask, request, jsonify, render_template
from database import init_db, get_connection
from gemini_service import generate_micro_modules
import PyPDF2
import docx
import io
import os

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max upload

# Initialize database on startup
init_db()

# ─────────────────────────────────────────────
# Helper: extract text from uploaded file
# ─────────────────────────────────────────────
def extract_text(file) -> str:
    filename = file.filename.lower()

    if filename.endswith(".pdf"):
        reader = PyPDF2.PdfReader(io.BytesIO(file.read()))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    elif filename.endswith(".docx"):
        doc = docx.Document(io.BytesIO(file.read()))
        return "\n".join(para.text for para in doc.paragraphs)

    else:  # plain text
        return file.read().decode("utf-8", errors="ignore")


# ─────────────────────────────────────────────
# Routes
# ─────────────────────────────────────────────
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/domains", methods=["GET"])
def get_domains():
    conn = get_connection()
    domains = conn.execute("SELECT domain_id, domain_name, description FROM KnowledgeDomains").fetchall()
    conn.close()
    return jsonify([dict(d) for d in domains])


@app.route("/api/upload", methods=["POST"])
def upload():
    # Validate file
    if "file" not in request.files:
        return jsonify({"error": "No file provided"}), 400

    file = request.files["file"]
    if file.filename == "":
        return jsonify({"error": "Empty filename"}), 400

    # Validate domains
    domain_ids = request.form.getlist("domain_ids")
    if not domain_ids:
        return jsonify({"error": "Please select at least one domain"}), 400

    # Extract text
    raw_text = extract_text(file)
    if not raw_text.strip():
        return jsonify({"error": "Could not extract text from file"}), 400

    # Get domain names for prompt
    conn = get_connection()
    placeholders = ",".join("?" * len(domain_ids))
    domain_rows = conn.execute(
        f"SELECT domain_id, domain_name FROM KnowledgeDomains WHERE domain_id IN ({placeholders})",
        domain_ids
    ).fetchall()
    domain_names = [r["domain_name"] for r in domain_rows]

    # Save source document
    cursor = conn.cursor()
    cursor.execute(
        "INSERT INTO SourceDocuments (file_name, raw_text) VALUES (?, ?)",
        (file.filename, raw_text)
    )
    doc_id = cursor.lastrowid

    # Save domain mappings
    for domain_id in domain_ids:
        cursor.execute(
            "INSERT INTO Document_Domain_Map (doc_id, domain_id) VALUES (?, ?)",
            (doc_id, domain_id)
        )

    conn.commit()

    # Generate micro-modules via Gemini
    modules = generate_micro_modules(raw_text, domain_names)

    # Save micro-modules
    for m in modules:
        cursor.execute(
            "INSERT INTO MicroModules (doc_id, module_title, module_content, reading_time_minutes) VALUES (?, ?, ?, ?)",
            (doc_id, m.get("title", "Module"), m.get("content", ""), m.get("reading_time_minutes", 2))
        )

    conn.commit()
    conn.close()

    return jsonify({
        "doc_id": doc_id,
        "file_name": file.filename,
        "domains": domain_names,
        "raw_text": raw_text,
        "modules": modules
    })


@app.route("/api/history", methods=["GET"])
def history():
    conn = get_connection()
    docs = conn.execute("""
        SELECT sd.doc_id, sd.file_name, sd.upload_timestamp,
               GROUP_CONCAT(kd.domain_name, ', ') as domains
        FROM SourceDocuments sd
        LEFT JOIN Document_Domain_Map ddm ON sd.doc_id = ddm.doc_id
        LEFT JOIN KnowledgeDomains kd ON ddm.domain_id = kd.domain_id
        GROUP BY sd.doc_id
        ORDER BY sd.upload_timestamp DESC
    """).fetchall()
    conn.close()
    return jsonify([dict(d) for d in docs])


@app.route("/api/modules/<int:doc_id>", methods=["GET"])
def get_modules(doc_id):
    conn = get_connection()
    modules = conn.execute(
        "SELECT * FROM MicroModules WHERE doc_id = ?", (doc_id,)
    ).fetchall()
    conn.close()
    return jsonify([dict(m) for m in modules])


if __name__ == "__main__":
    app.run(debug=True)
